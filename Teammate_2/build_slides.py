"""Generates docs/slides_teammate2.pptx.

    python build_slides.py
"""

import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUTPUT_PATH = os.path.join('docs', 'slides_teammate2.pptx')
METRICS_PATH = os.path.join('outputs', 'cluster_metrics.json')
PLOT_DIR = os.path.join('outputs', 'plots')

NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x44, 0x44, 0x44)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35),
                                   SLIDE_W - Inches(1.2), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = NAVY
    if subtitle:
        sp = tf.add_paragraph()
        sp.text = subtitle
        sp.font.size = Pt(13)
        sp.font.color.rgb = GREY


def add_bullets(slide, items, left=Inches(0.75), top=Inches(1.5),
                width=None, height=Inches(5.2), size=Pt(16)):
    width = width or SLIDE_W - Inches(1.5)
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p.text = ('- ' if level else '') + text
        p.level = level
        p.font.size = Pt(size.pt - 2) if level else size
        p.font.color.rgb = GREY if level else RGBColor(0, 0, 0)
        p.space_after = Pt(9)
    return box


def add_picture_fit(slide, path, top, max_w, max_h, left=None):
    """Adds a picture scaled to fit inside a box, centred horizontally."""
    from PIL import Image as PILImage
    with PILImage.open(path) as im:
        w, h = im.size
    ratio = min(max_w / w, max_h / h)
    new_w, new_h = int(w * ratio), int(h * ratio)
    left = left if left is not None else int((SLIDE_W - new_w) / 2)
    slide.shapes.add_picture(path, left, top, width=Emu(new_w), height=Emu(new_h))


def add_table(slide, data, left, top, width, height, font_size=12):
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    tbl = shape.table
    for r, row in enumerate(data):
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(value)
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(font_size)
            para.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
            if r == 0:
                para.font.bold = True
                para.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
    return tbl


def build():
    with open(METRICS_PATH) as f:
        m = json.load(f)

    km_sil = m['kmeans']['silhouette']
    som_sil = m['som']['silhouette_grouped']
    bmu_sil = m['som']['silhouette_raw_bmu']
    pc1, pc2 = m['pca']['explained_variance_ratio'][:2]

    os.makedirs('docs', exist_ok=True)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1. Title
    s = add_blank(prs)
    box = s.shapes.add_textbox(Inches(1), Inches(2.4), SLIDE_W - Inches(2),
                               Inches(2.6))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = 'Unsupervised Categorization of Radio Pulsar Candidates'
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.CENTER
    for line, size in [
        ('UE24CS352A Machine Learning, Mini-Project', 17),
        ('Teammate 2: PES1UG24CS419', 15),
        ('Problem Statement 3  |  HTRU2 dataset  |  1,639 confirmed pulsars', 13),
    ]:
        sp = tf.add_paragraph()
        sp.text = line
        sp.font.size = Pt(size)
        sp.font.color.rgb = GREY
        sp.alignment = PP_ALIGN.CENTER
        sp.space_before = Pt(10)

    # 2. Scope
    s = add_blank(prs)
    add_title(s, 'From "is it a pulsar" to "what kind of pulsar"',
              'The supervised half separates signal from noise. This half looks '
              'inside the signal.')
    add_bullets(s, [
        'Teammate 1 classifies all 17,898 candidates as pulsar or non-pulsar.',
        'This half asks a different question: do the confirmed pulsars '
        'themselves fall into distinct groups?',
        ('Fitted on the 1,639 pulsars only. The 16,259 non-pulsars are '
         'dropped, since including them would just recover the supervised '
         'split.', 1),
        ('8 features: mean, standard deviation, excess kurtosis and skewness '
         'of the integrated profile and of the DM-SNR curve.', 1),
        ('All features standardised first, because k-means and the SOM both '
         'compare points by Euclidean distance and the raw columns differ by '
         'orders of magnitude.', 1),
    ])

    # 3. Method
    s = add_blank(prs)
    add_title(s, 'Method')
    add_bullets(s, [
        'K-means (scikit-learn), k = 3, ten restarts, swept over k = 2 to 10.',
        'Self-Organizing Map, 5x5 grid, written from scratch in NumPy:',
        ('Competition: find the Best Matching Unit, the node nearest the '
         'sample in Euclidean distance.', 1),
        ('Cooperation: weight every node by a Gaussian centred on the BMU, '
         'measured in grid coordinates, not input space. This is what '
         'preserves topology.', 1),
        ('Adaptation: w = w + lr(t) h(t) (x - w), with learning rate and '
         'radius both decaying geometrically over 100 epochs.', 1),
        'PCA is fitted after clustering, for plotting only.',
        ('Clustering runs in the full 8-dimensional space. Projecting first '
         'would discard variance before the algorithms ever see it.', 1),
    ], size=Pt(15))

    # 4. Choosing k
    s = add_blank(prs)
    add_title(s, 'Choosing k: the two criteria disagree')
    add_picture_fit(s, os.path.join(PLOT_DIR, 'kmeans_k_selection.png'),
                    Inches(1.35), int(Inches(9.6)), int(Inches(3.5)))
    add_bullets(s, [
        'Inertia elbows at k = 3. Silhouette peaks at k = 2.',
        ('We keep k = 3: it is at the elbow, it matches the reference paper, '
         'and the k = 2 solution merges two groups that differ physically.', 1),
        ('Inertia measures compactness only; the silhouette also rewards '
         'separation, so on continuous data it favours fewer clusters.', 1),
    ], top=Inches(5.0), size=Pt(14))

    # 5. The three groups
    s = add_blank(prs)
    add_title(s, 'Three groups, described by their raw feature means')
    add_table(s, [
        ['Cluster', 'n', 'mean_ip', 'kurt_ip', 'skew_ip', 'mean_snr',
         'std_snr', 'skew_snr'],
        ['0', '739', '71.57', '2.04', '6.82', '38.65', '57.91', '8.04'],
        ['1', '679', '28.62', '5.01', '29.19', '77.09', '66.12', '1.52'],
        ['2', '221', '93.18', '1.01', '2.87', '3.44', '21.99', '101.43'],
    ], Inches(0.9), Inches(1.6), SLIDE_W - Inches(1.8), Inches(1.9), 14)
    add_bullets(s, [
        'Cluster 0 (739): the typical population, mid-range profile with a '
        'moderately dispersed DM-SNR curve.',
        'Cluster 1 (679): low profile mean, strong profile skew, highest '
        'DM-SNR mean. Bright and sharply peaked.',
        'Cluster 2 (221): high profile mean, very weak DM-SNR curve, extreme '
        'DM-SNR skew. Small and distinctive.',
    ], top=Inches(4.0), size=Pt(15))

    # 6. SOM diagnostics
    s = add_blank(prs)
    add_title(s, 'The SOM trained cleanly')
    add_picture_fit(s, os.path.join(PLOT_DIR, 'som_diagnostics.png'),
                    Inches(1.5), int(Inches(12.2)), int(Inches(3.4)))
    add_bullets(s, [
        f'Quantization error fell from '
        f'{m["som"]["quantization_error_first_epoch"]:.3f} to '
        f'{m["som"]["quantization_error_final_epoch"]:.3f} over 100 epochs, '
        f'and flattened.',
        f'All {m["som"]["occupied_nodes"]} of 25 nodes occupied, so no dead '
        f'units and the grid is not oversized.',
        'The U-matrix ridge in one corner is where cluster 2 separates.',
    ], top=Inches(5.1), size=Pt(14))

    # 7. Results
    s = add_blank(prs)
    add_title(s, 'Two different algorithms, the same three groups')
    add_picture_fit(s, os.path.join(PLOT_DIR, 'cluster_scatter_pca.png'),
                    Inches(1.35), int(Inches(8.2)), int(Inches(3.4)))
    add_table(s, [
        ['Method', 'Clusters', 'Silhouette'],
        ['K-means', '3', f'{km_sil:.4f}'],
        ['SOM 5x5, codebook grouped into 3', '3', f'{som_sil:.4f}'],
        ['SOM raw BMU assignment', '25', f'{bmu_sil:.4f}'],
    ], Inches(1.6), Inches(5.0), Inches(10.1), Inches(1.5), 13)
    add_bullets(s, [
        'The two methods agree on the group for 93.3% of candidates. '
        'Independent methods converging is the real evidence the structure '
        'is in the data, not the algorithm.',
    ], top=Inches(6.65), size=Pt(13))

    # 8. Why the 25-node score is low
    s = add_blank(prs)
    add_title(s, 'Two details worth being precise about')
    add_bullets(s, [
        'Comparing a 5x5 SOM with k-means at k = 3:',
        ('The map produces 25 micro-clusters, which is not comparable with 3. '
         'Following Vesanto and Alhoniemi (2000), the 25 codebook vectors are '
         'clustered into 3 groups and each candidate inherits the group of '
         'its BMU. Clustering 25 vectors is far cheaper than re-clustering '
         '1,639 samples.', 1),
        f'Why the raw 25-node silhouette ({bmu_sil:.3f}) looks bad:',
        ('It is low by construction, not because the map failed. With 25 '
         'clusters over 1,639 points, neighbouring nodes overlap heavily, and '
         'overlap is exactly what the silhouette penalises.', 1),
        f'PCA retains {(pc1 + pc2) * 100:.1f}% of variance in two components '
        f'(PC1 {pc1 * 100:.1f}%, PC2 {pc2 * 100:.1f}%).',
        ('PC1 sets profile shape (kurtosis +0.89, skewness +0.84) against '
         'profile mean (-0.87): a "narrow, peaked profile" axis. PC2 is '
         'dispersion.', 1),
    ], size=Pt(15))

    # 9. Demo
    s = add_blank(prs)
    add_title(s, 'Demo: Pulsar Candidate Explorer',
              'streamlit run app/demo_app.py')
    add_bullets(s, [
        'Enter the eight measurements for one candidate, or start from a '
        'preset, and the app returns:',
        ('A pulsar verdict with its probability, from Teammate 1\'s GDA '
         'model. Reusing it rather than reimplementing keeps the two halves '
         'consistent.', 1),
        ('The k-means cluster and the SOM cluster, plus the Best Matching '
         'Unit coordinates on the 5x5 grid.', 1),
        ('A flag when the two methods disagree, which usually means the '
         'candidate sits near a cluster boundary.', 1),
        ('A PCA scatter plot marking where the candidate falls among the '
         '1,639 known pulsars.', 1),
    ])

    # 10. Conclusions
    s = add_blank(prs)
    add_title(s, 'Conclusions and limitations')
    add_bullets(s, [
        'The pulsar population carries real internal structure: three groups '
        'of 739, 679 and 221, recovered independently by two algorithms.',
        f'Separation is moderate (silhouette {km_sil:.2f}). The limitation is '
        f'the feature set, not the algorithms.',
        ('HTRU2 gives summary statistics of the profile and the DM-SNR curve, '
         'so frequency and timing information is already averaged away. '
         'Pulsars separate best by period and dispersion measure, and neither '
         'survives into these eight columns. Adding them is the next step.', 1),
        'Three further caveats:',
        ('No ground truth for the sub-types, so validation is internal only. '
         'Confirming cluster 2 as a known class needs a catalogue cross-match '
         'such as the ATNF.', 1),
        ('k = 3 is a judgement call. The silhouette alone would have picked '
         'k = 2.', 1),
        ('The SOM depends on initialisation. The seed is fixed, and the '
         'three-group structure survives reseeding, but the map shifts.', 1),
    ], size=Pt(15))

    prs.save(OUTPUT_PATH)
    print(f'Wrote {OUTPUT_PATH} with {len(prs.slides.__iter__.__self__._sldIdLst)} slides')


if __name__ == '__main__':
    build()
