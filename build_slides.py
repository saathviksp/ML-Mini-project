import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_pptx_slides(output_path='docs/slides_teammate1.pptx'):
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    NAVY = RGBColor(26, 54, 93)
    DARK_BLUE = RGBColor(43, 108, 176)
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(113, 128, 150)
    DARK_TEXT = RGBColor(45, 55, 72)
    ACCENT = RGBColor(221, 107, 32)
    CARD_BG = RGBColor(247, 250, 252)

    def add_header(slide, title_text, category_text="SUPERVISED PIPELINE (TEAMMATE 1)"):
        # Header Box
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
        tf = header_box.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = ACCENT
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

    # Slide 1: Title Slide
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY
    
    title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.5))
    tf1 = title_box.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "Pulsar Candidate Classification"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    p2 = tf1.add_paragraph()
    p2.text = "UE24CS352A Machine Learning Mini-Project • Teammate 1 (Supervised Half)"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.space_before = Pt(10)
    
    p3 = tf1.add_paragraph()
    p3.text = "GDA Baseline (Scratch) | Random Forest (Scikit-Learn & Scratch) | Fold-Level Upsampling & Threshold Tuning"
    p3.font.size = Pt(14)
    p3.font.color.rgb = ACCENT
    p3.space_before = Pt(20)

    # Slide 2: Problem Statement & HTRU2 Dataset
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "Problem Statement & HTRU2 Dataset Overview")
    
    c1 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.0))
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Astronomical Context & Challenge"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = DARK_BLUE
    
    bullets = [
        "Radio pulsar search surveys generate millions of candidate signals.",
        "The vast majority are spurious noise or Radio Frequency Interference (RFI).",
        "HTRU2 Dataset: 17,898 total candidates collected by HTRUS survey.",
        "Severe Class Imbalance: 16,259 non-pulsars (90.84%) vs 1,639 true pulsars (9.16%). Imbalance ratio ~1:9.9.",
        "Goal: Automated binary classification to isolate true pulsars with high recall."
    ]
    for b in bullets:
        bp = tf.add_paragraph()
        bp.text = "• " + b
        bp.font.size = Pt(13)
        bp.font.color.rgb = DARK_TEXT
        bp.space_before = Pt(8)

    if os.path.exists('outputs/plots/class_distribution.png'):
        slide2.shapes.add_picture('outputs/plots/class_distribution.png', Inches(7.0), Inches(1.8), width=Inches(5.5))

    # Slide 3: Supervised Methodology & Equations
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Supervised Learning Models & Mathematical Formulations")
    
    tb = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "1. Gaussian Discriminant Analysis (GDA Baseline from Scratch)"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = DARK_BLUE
    
    gda_pts = [
        "Generative model assuming class-conditional Gaussian distributions: y ~ Bernoulli(φ), x|y=k ~ N(μ_k, Σ).",
        "Prior: φ = (1/N) Σ 1{y^(i) = 1}",
        "Means: μ_0 = Σ 1{y^(i)=0} x^(i) / Σ 1{y^(i)=0},   μ_1 = Σ 1{y^(i)=1} x^(i) / Σ 1{y^(i)=1}",
        "Shared Covariance: Σ = (1/N) Σ (x^(i) - μ_y^(i))(x^(i) - μ_y^(i))^T",
        "Posterior Probability: P(y=1|x) computed via sigmoid of log-likelihood ratio."
    ]
    for pt in gda_pts:
        bp = tf.add_paragraph()
        bp.text = "  • " + pt
        bp.font.size = Pt(13)
        bp.font.color.rgb = DARK_TEXT
        bp.space_before = Pt(4)

    p2 = tf.add_paragraph()
    p2.text = "2. Random Forest (Scikit-Learn & From Scratch)"
    p2.font.bold = True
    p2.font.size = Pt(16)
    p2.font.color.rgb = DARK_BLUE
    p2.space_before = Pt(14)
    
    rf_pts = [
        "Ensemble of uncorrelated decision trees using Bootstrap Sampling and Random Feature Subsets (√p features).",
        "Splits evaluated by maximizing Information Gain: IG = E(parent) - (N_left/N) E(left) - (N_right/N) E(right).",
        "Shannon Entropy: E(S) = - Σ p(x) log_2 p(x)."
    ]
    for pt in rf_pts:
        bp = tf.add_paragraph()
        bp.text = "  • " + pt
        bp.font.size = Pt(13)
        bp.font.color.rgb = DARK_TEXT
        bp.space_before = Pt(4)

    # Slide 4: 5-Fold Cross Validation Protocol
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "5-Fold CV Protocol with Fold-Level Upsampling")
    
    c1 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = c1.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Preventing Data Leakage via Strict Fold Isolation"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = DARK_BLUE
    
    cv_pts = [
        "Data Split: 80% Stratified Training Set (N=14,318) and 20% Held-out Test Set (N=3,580).",
        "Class Imbalance Handling: Random Minority Oversampling (ROS) applied to training folds ONLY.",
        "Crucial Requirement: Validation folds and held-out test set are NEVER upsampled (kept at raw 1:9.9 distribution).",
        "Hyperparameter Search: Scikit-learn Random Forest tuned using RandomizedSearchCV optimizing for Recall.",
        "Tuned Hyperparameters: n_estimators=200, min_samples_split=2, max_features='log2', max_depth=110."
    ]
    for pt in cv_pts:
        bp = tf.add_paragraph()
        bp.text = "• " + pt
        bp.font.size = Pt(14)
        bp.font.color.rgb = DARK_TEXT
        bp.space_before = Pt(10)

    # Slide 5: Performance Comparison Table
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Empirical Performance Comparison (Test Set, N=3,580)")
    
    rows = 7
    cols = 7
    table_shape = slide5.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    table = table_shape.table
    
    headers = ["Model", "Threshold", "Accuracy", "Recall", "Precision", "Specificity", "ROC AUC"]
    data = [
        ["GDA Baseline (Scratch)", "0.50", "97.91%", "89.63%", "87.76%", "98.74%", "0.9708"],
        ["GDA Baseline (Opt Th.)", "0.39", "97.57%", "90.85%", "83.94%", "98.25%", "0.9708"],
        ["RF (Scikit-Learn)", "0.50", "97.88%", "88.41%", "88.41%", "98.83%", "0.9678"],
        ["RF (Scikit-Learn Tuned)", "0.20", "96.76%", "91.16%", "77.46%", "97.32%", "0.9678"],
        ["RF From Scratch", "0.50", "97.21%", "90.85%", "80.98%", "97.85%", "0.9752"],
        ["RF From Scratch (Th 0.2)", "0.20", "90.95%", "93.29%", "50.33%", "90.71%", "0.9752"]
    ]
    
    for col_idx, h in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.color.rgb = WHITE
            paragraph.font.size = Pt(13)
            paragraph.alignment = PP_ALIGN.CENTER
            
    for row_idx, r_data in enumerate(data):
        for col_idx, val in enumerate(r_data):
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = val
            cell.fill.solid()
            if row_idx % 2 == 0:
                cell.fill.fore_color.rgb = CARD_BG
            else:
                cell.fill.fore_color.rgb = WHITE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(12)
                paragraph.font.color.rgb = DARK_TEXT
                paragraph.alignment = PP_ALIGN.CENTER

    # Slide 6: Threshold Sweep Optimization
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Threshold Optimization (0.5 → 0.2 to Maximize Recall)")
    
    c1 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.0))
    tf = c1.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Why Lower Decision Threshold?"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = DARK_BLUE
    
    sweep_pts = [
        "Astronomical Tradeoff: Discarding a true pulsar star is far more costly than inspecting a false candidate.",
        "Benchmark Alignment: Debesai et al. lowered threshold from 0.5 to 0.2 to boost recall.",
        "Threshold Sweep Plot: Sensitivity (TPR), 1-Specificity (FPR), and TPR-FPR difference plotted from 0.0 to 1.0.",
        "Result: Moving threshold to ~0.2 raises recall to 91.16% (sklearn RF) and 93.29% (scratch RF)."
    ]
    for pt in sweep_pts:
        bp = tf.add_paragraph()
        bp.text = "• " + pt
        bp.font.size = Pt(13)
        bp.font.color.rgb = DARK_TEXT
        bp.space_before = Pt(8)

    if os.path.exists('outputs/plots/rf_threshold_sweep.png'):
        slide6.shapes.add_picture('outputs/plots/rf_threshold_sweep.png', Inches(6.8), Inches(1.8), width=Inches(5.7))

    # Slide 7: ROC and PRC Curves
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "ROC and Precision-Recall Diagnostic Curves")
    
    if os.path.exists('outputs/plots/roc_curve.png') and os.path.exists('outputs/plots/prc_curve.png'):
        slide7.shapes.add_picture('outputs/plots/roc_curve.png', Inches(0.8), Inches(1.8), width=Inches(5.7))
        slide7.shapes.add_picture('outputs/plots/prc_curve.png', Inches(6.8), Inches(1.8), width=Inches(5.7))

    # Slide 8: Conclusions & Teammate 2 Integration
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Summary & Handoff to Teammate 2 (Unsupervised)")
    
    c1 = slide8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = c1.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "Key Takeaways (Teammate 1)"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = DARK_BLUE
    
    summary_pts = [
        "Successfully implemented GDA Baseline from scratch matching mathematical formulas.",
        "Tuned Scikit-Learn Random Forest and implemented pure NumPy Random Forest from scratch.",
        "Rigorous 5-fold CV with fold-level upsampling prevented data leakage and produced realistic evaluation.",
        "Threshold sweep demonstrated optimal operational point at ~0.2 threshold maximizing pulsar recall.",
        "Handoff to Teammate 2: Positive candidate subset (N=1,639 pulsars) is structured and ready for PCA, K-means (k=3), and Self-Organizing Map (SOM 5x5 grid) unsupervised clustering."
    ]
    for pt in summary_pts:
        bp = tf.add_paragraph()
        bp.text = "• " + pt
        bp.font.size = Pt(14)
        bp.font.color.rgb = DARK_TEXT
        bp.space_before = Pt(10)

    prs.save(output_path)
    print(f"Presentation slides saved to {output_path}")

if __name__ == '__main__':
    create_pptx_slides()
